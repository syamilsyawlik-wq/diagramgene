"""
xlsx_to_csv.py
Konversi template TRM (.xlsx, sheet "Aplikasi") ke CSV format yang dibaca trm_mode.py.

Aturan:
  - Baris header dicari otomatis (baris yang memuat "Nama Aplikasi"); baris di bawahnya = sub-header.
  - Header 2 tingkat diratakan jadi "Grup - Subkolom"; grup diambil dari sel merge di atasnya.
  - Baris contoh ("(Contoh) ...") dan baris tanpa Nama Aplikasi dibuang.
  - Tanggal -> DD/MM/YYYY; newline di dalam sel (Alt+Enter) -> ", ".
  - Output: UTF-8 dengan BOM, delimiter koma, CRLF (sama dengan "Save As CSV UTF-8" Excel).

CLI:  python xlsx_to_csv.py input.xlsx output.csv
"""
import csv
import datetime as dt
import io
import re
import sys

from openpyxl import load_workbook

SHEET_NAME = "Aplikasi"
COL_NAME = "Nama Aplikasi"
EXAMPLE_RE = re.compile(r"^\s*\(contoh\)", re.I)


def _clean_header(v):
    return re.sub(r"\s*\n\s*", " ", str(v)).strip() if v is not None else ""


def _cell_text(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).replace("\r\n", "\n").strip()
    return re.sub(r"\s*\n\s*", ", ", s)


def _merged_value(ws, row, col):
    """Nilai sel, atau nilai sel kiri-atas kalau sel ini bagian dari range merge."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return ws.cell(row, col).value


def xlsx_to_records(data):
    wb = load_workbook(io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else data,
                       data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"Sheet '{SHEET_NAME}' tidak ditemukan. Sheet yang ada: {', '.join(wb.sheetnames)}")
    ws = wb[SHEET_NAME]

    hdr_row = next((r for r in range(1, 11)
                    if any(_clean_header(c.value) == COL_NAME for c in ws[r])), None)
    if hdr_row is None:
        raise ValueError(f"Kolom '{COL_NAME}' tidak ditemukan di 10 baris pertama sheet '{SHEET_NAME}'.")
    sub_row = hdr_row + 1

    header, last = [], 0
    for c in range(1, ws.max_column + 1):
        group = _clean_header(_merged_value(ws, hdr_row, c))
        sub = _clean_header(ws.cell(sub_row, c).value)
        # sub-header asli hanya kalau sel group TIDAK merge vertikal ke baris sub-header
        if sub and sub != group:
            name = f"{group} - {sub}" if group else sub
        else:
            name = group
        header.append(name)
        if name:
            last = c
    header = header[:last]
    name_idx = header.index(COL_NAME)

    records, skipped = [], []
    for r in range(sub_row + 1, ws.max_row + 1):
        row = [_cell_text(ws.cell(r, c).value) for c in range(1, last + 1)]
        name = row[name_idx]
        if not name:
            continue
        if EXAMPLE_RE.match(name):
            skipped.append(name)
            continue
        records.append(row)
    return header, records, skipped


def xlsx_to_csv_bytes(data):
    header, records, skipped = xlsx_to_records(data)
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(header)
    w.writerows(records)
    return buf.getvalue().encode("utf-8-sig"), {"n_rows": len(records), "skipped_examples": skipped}


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Pakai: python xlsx_to_csv.py input.xlsx output.csv")
    out, info = xlsx_to_csv_bytes(open(sys.argv[1], "rb").read())
    open(sys.argv[2], "wb").write(out)
    print(f"{info['n_rows']} aplikasi ditulis; baris contoh dibuang: {len(info['skipped_examples'])}")
