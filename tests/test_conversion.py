"""Contoh sintetis untuk melindungi kontrak ketiga mode konversi."""
import base64
import html
import io
from pathlib import Path
import unittest
import urllib.parse
import zlib

from openpyxl import Workbook

from drawio_to_sparx import convert, list_pages
from generic_mode import convert_generic
from preview import render_preview
from trm_mode import convert_trm
from xlsx_to_csv import xlsx_to_csv_bytes

ROOT = Path(__file__).resolve().parents[1]
DRAWIO = (ROOT / "examples/proses_sederhana.drawio").read_bytes()
CSV = (ROOT / "examples/portofolio_aplikasi.csv").read_bytes()


class ConversionTests(unittest.TestCase):
    def test_bpmn_layouts_preserve_process(self):
        for orient, keep in [("vertical", False), ("horizontal", False), ("vertical", True)]:
            with self.subTest(orient=orient, keep_layout=keep):
                js, pages, notes = convert(DRAWIO, orient=orient, keep_layout=keep)
                self.assertEqual(notes, [])
                self.assertEqual(len(pages), 1)
                pg = pages[0]
                self.assertEqual({n["kind"] for n in pg["nodes"]}, {"start", "task", "end"})
                self.assertEqual([(f["src"], f["tgt"]) for f in pg["flows"]],
                                 [("start", "task"), ("task", "end")])
                self.assertTrue(render_preview(pg).startswith(b"\x89PNG\r\n\x1a\n"))
                self.assertIn('"Periksa permintaan"', js)
                self.assertNotIn("__DATA__", js)

    def test_generic_preserves_relative_geometry_and_labels(self):
        js, pages, notes = convert_generic(DRAWIO)
        self.assertEqual(notes, [])
        items = {it["key"]: it for it in pages[0]["items"]}
        self.assertEqual(len(items), 3)
        self.assertEqual(items["task"]["text"], "Periksa permintaan")
        self.assertEqual(items["task"]["r"] - items["task"]["l"], 120)
        self.assertEqual(items["task"]["t"] - items["start"]["t"], 100)
        self.assertEqual(len(pages[0]["links"]), 2)
        self.assertTrue(render_preview(pages[0]).startswith(b"\x89PNG"))
        self.assertNotIn("__DATA__", js)

    def test_compressed_drawio_and_embedded_svg(self):
        xml = DRAWIO.split(b"<mxGraphModel>", 1)[1].split(b"</mxGraphModel>", 1)[0]
        xml = b"<mxGraphModel>" + xml + b"</mxGraphModel>"
        compressor = zlib.compressobj(wbits=-15)
        raw = urllib.parse.quote(xml.decode(), safe="").encode()
        packed = base64.b64encode(compressor.compress(raw) + compressor.flush()).decode()
        compressed = f'<mxfile><diagram name="Test">{packed}</diagram></mxfile>'
        svg = ('<svg xmlns="http://www.w3.org/2000/svg" content="' +
               html.escape(compressed, quote=True) + '"/>')
        for data in [compressed.encode(), svg.encode()]:
            with self.subTest(format=data[:8]):
                self.assertEqual(list_pages(data)[0][0], "Test")
                self.assertEqual(len(convert(data)[1][0]["nodes"]), 3)

    def test_page_selection(self):
        js, pages, _ = convert(DRAWIO, page_names=["Persetujuan"])
        self.assertEqual([pg["name"] for pg in pages], ["Persetujuan"])
        with self.assertRaisesRegex(ValueError, "Tidak ada halaman"):
            convert(DRAWIO, page_names=["Unknown"])

    def test_trm_maps_apps_and_shared_technology(self):
        js, model, stats, warnings = convert_trm(CSV)
        self.assertEqual(warnings, [])
        self.assertEqual(stats["n_apps"], 2)
        self.assertEqual(stats["n_domains"], 8)
        self.assertEqual(stats["n_components"], 81)
        self.assertEqual(stats["n_realisation"], 2)
        self.assertEqual(stats["n_tech"], 1)
        self.assertEqual(stats["n_dependency"], 2)
        redundant = [c for d in model["domains"] for c in d["components"] if len(c["apps"]) >= 2]
        self.assertEqual(len(redundant), 1)
        self.assertEqual(redundant[0]["name"], "D25. Work Order Management")
        self.assertNotIn("__DATA__", js)

    def test_trm_duplicate_names_do_not_collide_with_existing_suffixes(self):
        data = "Nama Aplikasi\nApp\nApp\nApp #2\nApp\nApp #3\n"
        _, model, stats, warnings = convert_trm(data)
        names = [a["name"] for a in model["apps"]]
        self.assertEqual(names, ["App", "App #4", "App #2", "App #5", "App #3"])
        self.assertEqual(len(set(names)), stats["n_apps"])
        self.assertEqual(len(warnings), 2)

    def test_trm_reports_every_application_without_actual_oda_mapping(self):
        data = ("Nama Aplikasi;ODA Domain;ODA Component\n"
                "Unmapped;;Custom component\n"
                "Empty;;\n"
                "Inferred;;D25. Work Order Management\n")
        _, model, stats, warnings = convert_trm(data)
        self.assertEqual(stats["unmapped_apps"], ["Unmapped", "Empty"])
        self.assertEqual(stats["n_realisation"], 1)
        mapped = {key for d in model["domains"] for c in d["components"] for key in c["apps"]}
        self.assertEqual(mapped, {model["apps"][2]["key"]})
        self.assertTrue(any("Unmapped" in w and "tanpa domain" in w for w in warnings))

    def test_excel_template_to_trm(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Aplikasi"
        ws["A1"] = "Nama Aplikasi"
        ws.merge_cells("A1:A2")
        ws["B1"] = "Deskripsi"
        ws.merge_cells("B1:B2")
        ws["C1"] = "Data Management"
        ws["C2"] = "Database"
        ws.append(["(Contoh) Demo", "Example", "MySQL"])
        ws.append(["App Excel", "Baris 1\nBaris 2", "PostgreSQL"])
        buf = io.BytesIO()
        wb.save(buf)
        csv, info = xlsx_to_csv_bytes(buf.getvalue())
        self.assertEqual(info["n_rows"], 1)
        self.assertEqual(len(info["skipped_examples"]), 1)
        _, model, stats, _ = convert_trm(csv)
        self.assertEqual(stats["n_apps"], 1)
        self.assertEqual(model["apps"][0]["notes"], "Baris 1, Baris 2")
        self.assertEqual(stats["n_tech"], 1)

    def test_invalid_inputs_report_errors(self):
        for data in [b"", b"Name\nApp A", b"Nama Aplikasi\n"]:
            with self.subTest(data=data), self.assertRaises(ValueError):
                convert_trm(data)
        with self.assertRaisesRegex(ValueError, "bukan diagram draw.io"):
            convert(b"<document/>")
        with self.assertRaisesRegex(ValueError, "tidak berisi data draw.io"):
            convert(b'<svg xmlns="http://www.w3.org/2000/svg"/>')


if __name__ == "__main__":
    unittest.main()
